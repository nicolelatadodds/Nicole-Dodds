import logging
import os
from datetime import date, datetime
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from models import SheetData, TaskRow

logger = logging.getLogger(__name__)

SHEET_NAME = "Client Status"

COLOR_HEADER_BG = "1F3864"
COLOR_HEADER_FG = "FFFFFF"
COLOR_ENDDATE_BG = "FFE699"
COLOR_ROW_STRIPE = "F2F2F2"


def _thin_border() -> Border:
    thin = Side(style="thin", color="CCCCCC")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _fmt_date(d: Optional[date]) -> str:
    if d is None:
        return "TBD"
    # %-d is Linux-only (no zero-pad); on Windows use %#d
    return d.strftime("%-d %b %Y")


def _bullet_text(row: TaskRow) -> str:
    return f"\u2022 {row.name} \u2014 {_fmt_date(row.start_date)} to {_fmt_date(row.end_date)}"


def write_client_status(sheet_data: SheetData, output_path: str) -> None:
    """
    Writes (or overwrites) the 'Client Status' sheet in the Excel file at
    output_path. Other sheets in an existing workbook are preserved.
    """
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

    if os.path.exists(output_path):
        wb = load_workbook(output_path)
        if SHEET_NAME in wb.sheetnames:
            del wb[SHEET_NAME]
        ws = wb.create_sheet(SHEET_NAME, 0)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_NAME

    ws.column_dimensions["A"].width = 85

    # Row 1: title
    ws["A1"] = "Client Status Report"
    ws["A1"].font = Font(name="Calibri", size=18, bold=True, color=COLOR_HEADER_FG)
    ws["A1"].fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 38

    # Row 2: last synced timestamp
    ws["A2"] = f"Last Synced: {datetime.now().strftime('%-d %b %Y  %-I:%M %p')}"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="666666")
    ws["A2"].alignment = Alignment(horizontal="right")

    # Row 3: spacer
    ws.row_dimensions[3].height = 8

    # Row 4: section label
    ws["A4"] = "PROJECT TIMING"
    ws["A4"].font = Font(name="Calibri", size=11, bold=True, color=COLOR_HEADER_BG)
    ws["A4"].alignment = Alignment(horizontal="left")

    # Rows 5+: bullet list
    start_row = 5
    for idx, task_row in enumerate(sheet_data.rows):
        r = start_row + idx
        cell = ws.cell(row=r, column=1, value=_bullet_text(task_row))
        cell.font = Font(name="Calibri", size=11)
        cell.alignment = Alignment(wrap_text=True, vertical="top", indent=2)
        cell.border = _thin_border()
        ws.row_dimensions[r].height = 20
        if idx % 2 == 1:
            cell.fill = PatternFill("solid", fgColor=COLOR_ROW_STRIPE)

    # Final end date block
    gap_row = start_row + len(sheet_data.rows) + 1
    label_cell = ws.cell(row=gap_row, column=1, value="FINAL END DATE")
    label_cell.font = Font(name="Calibri", size=11, bold=True, color=COLOR_HEADER_BG)
    label_cell.alignment = Alignment(horizontal="left")

    final_row = gap_row + 1
    final = sheet_data.final_end_date
    final_cell = ws.cell(
        row=final_row,
        column=1,
        value=_fmt_date(final) if final else "No end date found",
    )
    final_cell.font = Font(name="Calibri", size=14, bold=True)
    final_cell.fill = PatternFill("solid", fgColor=COLOR_ENDDATE_BG)
    final_cell.alignment = Alignment(horizontal="center", vertical="center")
    final_cell.border = _thin_border()
    ws.row_dimensions[final_row].height = 30

    wb.save(output_path)
    logger.info("Wrote Client Status sheet to %s", output_path)
